from utils.clock import now_tz
import asyncio
import os
import re
import tempfile
import uuid
import openpyxl
from typing import List, Dict, Any, Optional, Tuple
from datetime import datetime, date, time

from database.db import Database, normalize_phone
from utils.time_validation import validate_schedule_times

WEEKDAY_COLUMNS = {
    "dushanba": 0,
    "seshanba": 1,
    "chorshanba": 2,
    "payshanba": 3,
    "juma": 4,
    "shanba": 5,
    "yakshanba": 6,
}

_TIME_RANGE_RE = re.compile(
    r"^\s*([0-9]{1,2})[:.]([0-9]{2})\s*[-–—]\s*([0-9]{1,2})[:.]([0-9]{2})\s*$"
)
_TIME_SINGLE_RE = re.compile(r"^\s*([0-9]{1,2})[:.]([0-9]{2})\s*$")
_TIME_RANGE_ERROR = "Vaqt oralig'i HH:MM-HH:MM ko'rinishida bo'lishi kerak"


def _parse_time_range(value: Any) -> Tuple[str, str]:
    """Accept normal text plus Excel time/datetime cells and normalize to HH:MM."""
    if isinstance(value, datetime):
        formatted = value.strftime("%H:%M")
        return formatted, formatted
    if isinstance(value, time):
        formatted = value.strftime("%H:%M")
        return formatted, formatted
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel may expose a time-only cell as a fraction of a day.
        if 0 <= float(value) < 1:
            total_minutes = round(float(value) * 24 * 60)
            total_minutes %= 24 * 60
            return f"{total_minutes // 60:02d}:{total_minutes % 60:02d}", ""
    text = str(value or "").strip()
    match = _TIME_RANGE_RE.fullmatch(text)
    if match:
        start_hour, start_minute, end_hour, end_minute = match.groups()
        return (
            f"{start_hour.zfill(2)}:{start_minute}",
            f"{end_hour.zfill(2)}:{end_minute}",
        )
    single = _TIME_SINGLE_RE.fullmatch(text)
    if single:
        hour, minute = single.groups()
        return f"{hour.zfill(2)}:{minute}", ""
    raise ValueError(_TIME_RANGE_ERROR)


def _excel_text(value: Any) -> str:
    """Turn common Excel scalar types into stable user-facing text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def _normalize_match_text(value: Any) -> str:
    text = _normalize_apostrophes(_excel_text(value)).casefold()
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _normalize_phone_cell(value: Any) -> Optional[str]:
    """Handle Excel numeric cells, leading 998, spaces, punctuation and +998."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = _excel_text(value)
    # Excel sometimes stores +998... as a formula-like/text value or scientific notation.
    text = text.replace("\u00a0", " ").strip()
    if re.fullmatch(r"[0-9]+(?:\.0+)?[Ee][+-]?[0-9]+", text, re.I):
        try:
            text = str(int(float(text)))
        except (TypeError, ValueError, OverflowError):
            pass
    return normalize_phone(text)


class ExcelService:
    @staticmethod
    def generate_template(output_path: Optional[str] = None) -> str:
        if not output_path:
            stamp = now_tz().strftime("%Y%m%d_%H%M%S_%f")
            filename = f"shablon_oquvchilar_{stamp}_{uuid.uuid4().hex[:8]}.xlsx"
            output_path = os.path.join(tempfile.gettempdir(), filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "O'quvchilar"
        headers = [
            "F.I.Sh", "Mahalla", "Ota telefoni", "Ona telefoni",
            "Dushanba", "Seshanba", "Chorshanba", "Payshanba",
            "Juma", "Shanba", "Yakshanba",
        ]
        ws.append(headers)
        ws.append([
            "Ali Valiyev", "Arpaariq", "+998901234567", "+998907654321",
            "09:00-11:00", "", "09:00-11:00", "", "09:00-11:00", "", ""
        ])
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)
        wb.save(output_path)
        return output_path

    @staticmethod
    async def parse_and_preview_excel(file_path: str, db: Database) -> Dict[str, Any]:
        headers, sheet_rows = await asyncio.to_thread(_read_excel_rows, file_path)

        header_map: Dict[str, int] = {}
        seen_weekdays: Dict[int, int] = {}
        duplicate_weekdays = set()
        for col_idx, header in enumerate(headers, 1):
            val = _normalize_match_text(header)
            day_num = WEEKDAY_COLUMNS.get(val)
            if day_num is not None:
                if day_num in seen_weekdays:
                    duplicate_weekdays.add(day_num)
                else:
                    seen_weekdays[day_num] = col_idx
                    header_map[f"day_{day_num}"] = col_idx
                continue
            if any(k in val for k in ["f.i.sh", "fish", "fio", "ism"]):
                header_map["full_name"] = col_idx
            elif any(k in val for k in ["mahalla", "qishloq", "maktab"]):
                header_map["village"] = col_idx
            elif val in {"ona telefoni", "ona telefon", "telefon 2", "qo'shimcha telefon", "qo'shimcha raqam"}:
                header_map["phone_2"] = col_idx
            elif any(k in val for k in ["ota telefoni", "ota telefon", "asosiy telefon", "telefon 1"]):
                if "phone_1" not in header_map:
                    header_map["phone_1"] = col_idx
            elif "telefon" in val and "phone_1" not in header_map:
                header_map["phone_1"] = col_idx
            elif any(k in val for k in ["ona", "qo'shimcha"]):
                header_map["phone_2"] = col_idx

        if duplicate_weekdays:
            errors = [
                {
                    "row": 1,
                    "full_name": "-",
                    "error": f"Hafta kuni sarlavhasi takrorlangan: {day_name.capitalize()}",
                }
                for day_name, day_num in WEEKDAY_COLUMNS.items()
                if day_num in duplicate_weekdays
            ]
            return {"new_count": 0, "update_count": 0, "error_count": len(errors), "errors": errors, "valid_records": []}

        idx_name = header_map.get("full_name", 1)
        idx_village = header_map.get("village", 2)
        idx_phone1 = header_map.get("phone_1", 3)
        idx_phone2 = header_map.get("phone_2", 4)
        valid_records = []
        errors = []

        conn = await db.get_connection()
        try:
            # Match village names/schools ourselves after normalization so Excel casing,
            # apostrophes and extra spaces do not cause otherwise valid rows to fail.
            villages = await db.get_villages()
            village_by_name: Dict[str, list] = {}
            village_by_school: Dict[str, list] = {}
            for village in villages:
                village_by_name.setdefault(_normalize_match_text(village.get("name")), []).append(village)
                school = _normalize_match_text(village.get("school_name") or "")
                if school:
                    village_by_school.setdefault(school, []).append(village)

            duplicate_keys = set()
            for row_idx, row_cells in enumerate(sheet_rows, start=2):
                if not any(value is not None and _excel_text(value) for value in row_cells):
                    continue

                def cell_value(index: int):
                    return row_cells[index - 1] if index <= len(row_cells) else None

                full_name_val = _normalize_apostrophes(_excel_text(cell_value(idx_name)))
                village_val = _normalize_apostrophes(_excel_text(cell_value(idx_village)))
                phone1_raw = _excel_text(cell_value(idx_phone1))
                phone2_raw = _excel_text(cell_value(idx_phone2))

                if not full_name_val:
                    errors.append({"row": row_idx, "full_name": "-", "error": "F.I.Sh ko'rsatilmagan"})
                    continue
                if not village_val:
                    errors.append({"row": row_idx, "full_name": full_name_val, "error": "Mahalla ko'rsatilmagan"})
                    continue

                norm_p1 = _normalize_phone_cell(cell_value(idx_phone1))
                if not norm_p1:
                    errors.append({"row": row_idx, "full_name": full_name_val, "error": f"Asosiy telefon noto'g'ri: {phone1_raw}"})
                    continue
                norm_p2 = _normalize_phone_cell(cell_value(idx_phone2))

                village_key = _normalize_match_text(village_val)
                candidates = village_by_name.get(village_key) or village_by_school.get(village_key) or []
                village_id = candidates[0]["id"] if len(candidates) == 1 else None

                schedules = []
                sched_error = None
                for day_name, day_num in WEEKDAY_COLUMNS.items():
                    col_idx = header_map.get(f"day_{day_num}")
                    if col_idx is None:
                        continue
                    time_val = cell_value(col_idx)
                    if time_val is None or not _excel_text(time_val):
                        continue
                    try:
                        st, et = _parse_time_range(time_val)
                        if not et:
                            raise ValueError(_TIME_RANGE_ERROR)
                        validate_schedule_times(st, et)
                        schedules.append({"weekday": day_num, "start_time": st, "end_time": et})
                    except ValueError:
                        sched_error = f"{day_name.capitalize()} soati noto'g'ri: {time_val}. HH:MM-HH:MM ko'rinishida kiriting."
                        break

                if sched_error:
                    errors.append({"row": row_idx, "full_name": full_name_val, "error": sched_error})
                    continue

                if not village_id:
                    if len(candidates) > 1:
                        candidate_names = [v.get("name") or v.get("school_name") for v in candidates]
                        error = f"Mahalla bir nechta mos keldi: {', '.join(candidate_names)}"
                    else:
                        candidates = await db.suggest_villages(village_val)
                        error = f"Mahalla aniq topilmadi: {village_val}"
                    errors.append({
                        "row": row_idx,
                        "full_name": full_name_val,
                        "error": error,
                        "candidates": candidates,
                    })
                    continue

                duplicate_key = (full_name_val.casefold(), norm_p1)
                if duplicate_key in duplicate_keys:
                    errors.append({
                        "row": row_idx,
                        "full_name": full_name_val,
                        "error": "Takroriy o'quvchi qatori; birinchi yozuv saqlanadi",
                    })
                    continue
                duplicate_keys.add(duplicate_key)

                valid_records.append({
                    "full_name": full_name_val,
                    "village_id": village_id,
                    "parent_phone_1": norm_p1,
                    "parent_phone_2": norm_p2,
                    "schedules": schedules,
                })

            new_count = 0
            update_count = 0
            for rec in valid_records:
                cur_ex = await conn.execute(
                    "SELECT id FROM students WHERE full_name = ? AND parent_phone_1 = ?",
                    (rec["full_name"], rec["parent_phone_1"]),
                )
                if await cur_ex.fetchone():
                    update_count += 1
                else:
                    new_count += 1

            return {
                "new_count": new_count,
                "update_count": update_count,
                "error_count": len(errors),
                "errors": errors,
                "valid_records": valid_records,
            }
        finally:
            await conn.close()

    @staticmethod
    async def export_students_to_excel(db: Database, output_path: Optional[str] = None) -> str:
        if not output_path:
            output_path = f"baza_oquvchilar_{now_tz().strftime('%Y%m%d_%H%M%S')}.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "O'quvchilar Bazasi"
        headers = [
            "F.I.Sh", "Mahalla", "Ota telefoni", "Ona telefoni",
            "Dushanba", "Seshanba", "Chorshanba", "Payshanba",
            "Juma", "Shanba", "Yakshanba",
        ]
        ws.append(headers)

        conn = await db.get_connection()
        try:
            cur_students = await conn.execute("""
                SELECT s.id, s.full_name, s.parent_phone_1, s.parent_phone_2, v.name as village_name
                FROM students s
                LEFT JOIN villages v ON s.village_id = v.id
                WHERE s.is_active = 1
                ORDER BY s.id ASC
            """)
            students = [dict(r) for r in await cur_students.fetchall()]
            for s in students:
                cur_sched = await conn.execute("""
                    SELECT weekday, start_time, end_time
                    FROM schedules
                    WHERE student_id = ?
                    ORDER BY weekday ASC
                """, (s["id"],))
                sched_rows = await cur_sched.fetchall()
                sched_map = {row["weekday"]: f"{row['start_time']}-{row['end_time']}" for row in sched_rows}
                ws.append([
                    _excel_safe_text(s["full_name"]),
                    _excel_safe_text(s.get("village_name") or ""),
                    _excel_safe_text(s.get("parent_phone_1") or ""),
                    _excel_safe_text(s.get("parent_phone_2") or ""),
                    sched_map.get(0, ""), sched_map.get(1, ""), sched_map.get(2, ""),
                    sched_map.get(3, ""), sched_map.get(4, ""), sched_map.get(5, ""), sched_map.get(6, ""),
                ])
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 14)
            wb.save(output_path)
            return output_path
        finally:
            await conn.close()


def _excel_safe_text(value: Any) -> str:
    text = str(value if value is not None else "")
    if text.lstrip().startswith("="):
        return "'" + text
    return text


def _normalize_apostrophes(text: str) -> str:
    return re.sub(r"[ʻʼ‘’`´]", "'", text)


def _read_excel_rows(file_path: str) -> Tuple[List[Any], List[Tuple[Any, ...]]]:
    """Load .xlsx and legacy .xls cell values off the event loop."""
    extension = os.path.splitext(file_path)[1].lower()
    if extension == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(file_path, on_demand=True)
        try:
            worksheet = workbook.sheet_by_index(0)
            rows = []
            for row_idx in range(worksheet.nrows):
                row = []
                for col_idx in range(worksheet.ncols):
                    cell = worksheet.cell(row_idx, col_idx)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_NUMBER and value.is_integer():
                        value = int(value)
                    row.append(value)
                rows.append(tuple(row))
            headers = list(rows[0]) if rows else []
            return headers, rows[1:]
        finally:
            workbook.release_resources()

    workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = list(next(rows, ()))
        return headers, list(rows)
    finally:
        workbook.close()
