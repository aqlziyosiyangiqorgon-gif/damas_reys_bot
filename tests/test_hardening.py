import asyncio
import time

import pytest
from aiogram.fsm.storage.base import StorageKey
from aiogram.methods import SendMessage
from aiogram.exceptions import TelegramNetworkError
import openpyxl
from unittest.mock import AsyncMock

from database.db import Database
from middlewares.sqlite_storage import SQLiteFSMStorage
from utils.text import escape_html
from utils.time import now_tashkent
from utils.telegram import safe_telegram_call
from services.excel_service import ExcelService
from services.scheduler_service import SchedulerService
from services.trip_builder import TripBuilder
from utils.time import now_tashkent
from handlers.admin import _edit_broadcast_message
from types import SimpleNamespace
from config import settings
from handlers.admin import (
    BroadcastStates,
    EditStudentStates,
    PENDING_BROADCASTS,
    BROADCAST_TTL_SECONDS,
    _purge_expired_broadcasts,
    cb_broadcast_cancel,
    cb_broadcast_select_audience,
    cb_broadcast_select_student,
    cb_admin_nav,
    cb_broadcast_send,
    process_broadcast_content,
)


@pytest.mark.asyncio
async def test_atomic_trip_and_item_claims(tmp_path):
    db = Database(str(tmp_path / "hardening.db"))
    await db.init_db()
    driver_id = await db.add_driver(10001, "Driver", "+998901234567")
    student_id = await db.add_student("Student", 1, "+998901234567")

    conn = await db.get_connection()
    cur = await conn.execute(
        "INSERT INTO trips (date, direction, driver_id, planned_time, status, approved) "
        "VALUES ('2026-10-14', 'to_center', ?, '08:00', 'approved', 1)",
        (driver_id,),
    )
    trip_id = cur.lastrowid
    cur_item = await conn.execute(
        "INSERT INTO trip_items (trip_id, student_id, village_id, status) VALUES (?, ?, 1, 'pending')",
        (trip_id, student_id),
    )
    item_id = cur_item.lastrowid
    await conn.commit()
    await conn.close()

    transitions = await asyncio.gather(*[
        db.transition_trip(
            trip_id, ["approved"], "on_way", started_at="2026-10-14 08:00:00",
            direction="to_center", actor_tg_id=10001, event_action="start_trip",
        )
        for _ in range(2)
    ])
    assert sorted(transitions) == [False, True]

    claims = await asyncio.gather(*[
        db.claim_item(
            item_id, trip_id, ["pending"], "picked_up", "picked_at", "2026-10-14 08:01:00",
            driver_id, trip_statuses=["on_way"], direction="to_center",
            actor_tg_id=10001, event_action="pickup_student",
        )
        for _ in range(2)
    ])
    assert sorted(claims) == [False, True]

    conn = await db.get_connection()
    cur = await conn.execute("SELECT COUNT(*) FROM event_log WHERE trip_id = ?", (trip_id,))
    assert (await cur.fetchone())[0] == 2
    await conn.close()


@pytest.mark.asyncio
async def test_fsm_storage_and_admin_workflows_persist(tmp_path):
    db = Database(str(tmp_path / "storage.db"))
    await db.init_db()
    key = StorageKey(bot_id=7, chat_id=8, user_id=9, destiny="default")

    first = SQLiteFSMStorage(db)
    await first.set_state(key, "AdminStates:waiting")
    await first.set_data(key, {"student_id": 42})
    second = SQLiteFSMStorage(db)
    assert await second.get_state(key) == "AdminStates:waiting"
    assert await second.update_data(key, {"name": "Ali"}) == {"student_id": 42, "name": "Ali"}
    assert await first.get_data(key) == {"student_id": 42, "name": "Ali"}

    payload = {"records": [{"full_name": "Ali", "village_id": 1}]}
    await db.save_workflow_payload("batch-1", 99, "excel", payload)
    assert await db.take_workflow_payload("batch-1", 100, "excel") is None
    assert await db.take_workflow_payload("batch-1", 99, "excel") == payload
    assert await db.take_workflow_payload("batch-1", 99, "excel") is None


def test_html_and_application_timezone_helpers():
    assert escape_html('<name "x">') == '&lt;name "x"&gt;'
    assert now_tashkent().utcoffset().total_seconds() == 5 * 60 * 60


@pytest.mark.asyncio
async def test_telegram_retry_handles_transient_errors():
    method = AsyncMock(side_effect=[TelegramNetworkError(SendMessage(chat_id=1, text="x"), "offline"), None])
    assert await safe_telegram_call(method, retries=2)
    assert method.await_count == 2


@pytest.mark.asyncio
async def test_excel_time_validation_and_formula_safe_export(tmp_path):
    db = Database(str(tmp_path / "excel.db"))
    await db.init_db()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["F.I.Sh", "Mahalla", "Ota telefoni", "Ona telefoni", "Dushanba"])
    sheet.append(["Student", "Arpaariq", "+998901234567", "", "10:00-09:00"])
    source = tmp_path / "invalid.xlsx"
    workbook.save(source)
    workbook.close()

    preview = await ExcelService.parse_and_preview_excel(str(source), db)
    assert preview["valid_records"] == []
    assert preview["error_count"] == 1
    assert "tugash vaqti" in preview["errors"][0]["error"]

    await db.add_student("=1+1", 1, "+998901234567")
    exported = await ExcelService.export_students_to_excel(db, str(tmp_path / "export.xlsx"))
    exported_workbook = openpyxl.load_workbook(exported, data_only=False)
    assert exported_workbook.active.cell(row=2, column=1).value == "'=1+1"
    assert exported_workbook.active.cell(row=2, column=1).data_type == "s"
    exported_workbook.close()


@pytest.mark.asyncio
async def test_excel_import_normalizes_names_matches_villages_exactly_and_maps_parent_phones(tmp_path):
    db = Database(str(tmp_path / "excel_normalization.db"))
    await db.init_db()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["F.I.Sh", "Mahalla", "Ona telefoni", "Ota telefoni"])
    sheet.append(["Po'latov", "Arpaariq", "+998907654321", "+998901234567"])
    sheet.append(["Po‘latov", "Arpaariq", "+998907654321", "+998901234567"])
    sheet.append(["Another Student", "Galaravat 3", "", "+998909876543"])
    source = tmp_path / "students.XLSX"
    workbook.save(source)
    workbook.close()

    preview = await ExcelService.parse_and_preview_excel(str(source), db)

    assert len(preview["valid_records"]) == 1
    student = preview["valid_records"][0]
    assert student["full_name"] == "Po'latov"
    assert student["parent_phone_1"] == "+998901234567"
    assert student["parent_phone_2"] == "+998907654321"
    unmatched = next(error for error in preview["errors"] if error["row"] == 4)
    assert "Mahalla aniq topilmadi" in unmatched["error"]
    assert "Galaravat" in unmatched["candidates"]


@pytest.mark.asyncio
async def test_student_driver_assignment_is_used_and_separates_trips(tmp_path):
    db = Database(str(tmp_path / "assignments.db"))
    await db.init_db()
    driver_a = await db.add_driver(20001, "Driver A", "+998901111111")
    driver_b = await db.add_driver(20002, "Driver B", "+998902222222")
    student_a = await db.add_student("Student A", 1, "+998903333333")
    student_b = await db.add_student("Student B", 1, "+998904444444")
    assert await db.set_student_driver(student_a, driver_a)
    assert await db.set_student_driver(student_b, driver_b)

    weekday = now_tashkent().weekday()
    await db.add_schedule(student_a, weekday, "09:00", "11:00")
    await db.add_schedule(student_b, weekday, "09:00", "11:00")
    rows = await db.get_day_schedules(weekday)
    assert {row["student_id"]: row["default_driver_id"] for row in rows} == {
        student_a: driver_a,
        student_b: driver_b,
    }

    trips, _ = TripBuilder().generate_trips_for_day(rows)
    assert len(trips) == 4
    assert {trip["driver_id"] for trip in trips} == {driver_a, driver_b}
    assert all(len(trip["items"]) == 1 for trip in trips)

    scheduler = SchedulerService(db)
    success, _, generated = await scheduler.run_daily_plan(now_tashkent().strftime("%Y-%m-%d"))
    assert success
    assert len(generated) == 4
    saved = await db.get_trips_by_date(now_tashkent().strftime("%Y-%m-%d"))
    assert len(saved) == 4
    assert all(trip["driver_name"] in {"Driver A", "Driver B"} for trip in saved)
    await db.approve_all_daily_trips(now_tashkent().strftime("%Y-%m-%d"), admin_id=1)
    success, _, _ = await scheduler.run_daily_plan(now_tashkent().strftime("%Y-%m-%d"))
    assert not success
    assert len(await db.get_trips_by_date(now_tashkent().strftime("%Y-%m-%d"))) == 4


@pytest.mark.asyncio
async def test_broadcast_status_edit_uses_caption_for_photo_preview():
    photo_message = SimpleNamespace(
        photo=[object()],
        edit_caption=AsyncMock(),
        edit_text=AsyncMock(),
    )
    assert await _edit_broadcast_message(photo_message, "sent")
    photo_message.edit_caption.assert_awaited_once_with(caption="sent", parse_mode="HTML")
    photo_message.edit_text.assert_not_awaited()

    text_message = SimpleNamespace(photo=None, edit_caption=AsyncMock(), edit_text=AsyncMock())
    assert await _edit_broadcast_message(text_message, "sent")
    text_message.edit_text.assert_awaited_once_with(text="sent", parse_mode="HTML")


@pytest.mark.asyncio
async def test_student_navigation_enters_edit_state(tmp_path):
    db = Database(str(tmp_path / "navigation.db"))
    await db.init_db()
    await db.add_student("Student", 1, "+998901234567")

    admin_id = settings.admin_ids_list[0]
    callback = SimpleNamespace(
        from_user=SimpleNamespace(id=admin_id),
        data="admin_nav:students",
        message=SimpleNamespace(edit_text=AsyncMock()),
        answer=AsyncMock(),
    )
    state = AsyncMock()
    await cb_admin_nav(callback, db, state)
    state.set_state.assert_awaited_once_with(EditStudentStates.selecting_student)
    assert "Faol o'quvchilar ro'yxati" in callback.message.edit_text.await_args.args[0]


@pytest.mark.asyncio
async def test_admin_can_reassign_only_unstarted_trip(tmp_path):
    db = Database(str(tmp_path / "trip-driver.db"))
    await db.init_db()
    driver_a = await db.add_driver(20004, "Driver A", "+998901234567")
    driver_b = await db.add_driver(20005, "Driver B", "+998902345678")
    conn = await db.get_connection()
    cur = await conn.execute(
        "INSERT INTO trips (date, direction, driver_id, planned_time, status, approved) "
        "VALUES ('2026-10-14', 'to_center', ?, '08:00', 'approved', 1)",
        (driver_a,),
    )
    trip_id = cur.lastrowid
    await conn.commit()
    await conn.close()

    assert await db.assign_driver_to_trip(trip_id, driver_b)
    assert (await db.get_trip_details(trip_id))["driver_id"] == driver_b
    conn = await db.get_connection()
    await conn.execute("UPDATE trips SET status = 'on_way' WHERE id = ?", (trip_id,))
    await conn.commit()
    await conn.close()
    assert not await db.assign_driver_to_trip(trip_id, driver_a)


@pytest.mark.asyncio
async def test_broadcast_plain_text_and_photo_status_are_safe(tmp_path):
    db = Database(str(tmp_path / "broadcast.db"))
    await db.init_db()
    driver_tg_id = 20003
    await db.add_driver(driver_tg_id, "Driver", "+998901234567")
    admin_id = settings.admin_ids_list[0]

    state = AsyncMock()
    state.get_data.return_value = {"bc_aud_type": "drivers", "bc_aud_val": None}
    source_message = SimpleNamespace(
        from_user=SimpleNamespace(id=admin_id),
        text="<xabar> & oddiy matn",
        caption=None,
        photo=None,
        answer=AsyncMock(),
    )
    await process_broadcast_content(source_message, state, db)
    preview = source_message.answer.await_args.args[0]
    assert "&lt;xabar&gt; &amp; oddiy matn" in preview

    batch_id = next(reversed(PENDING_BROADCASTS))
    outgoing_bot = AsyncMock()
    callback = SimpleNamespace(
        from_user=SimpleNamespace(id=admin_id),
        data=f"bc_send:{batch_id}",
        message=SimpleNamespace(photo=[object()], edit_caption=AsyncMock(), edit_text=AsyncMock()),
        answer=AsyncMock(),
        bot=outgoing_bot,
    )
    await cb_broadcast_send(callback, db)
    assert callback.answer.await_count == 1
    outgoing_bot.send_message.assert_awaited_once_with(
        chat_id=driver_tg_id, text="<xabar> & oddiy matn", parse_mode=None
    )
    callback.message.edit_caption.assert_awaited()


@pytest.mark.asyncio
async def test_broadcast_student_audience_uses_exact_student_id(tmp_path):
    db = Database(str(tmp_path / "student_audience.db"))
    await db.init_db()
    first_student = await db.add_student("Ali Valiyev", 1, "+998901111111")
    second_student = await db.add_student("Ali Valiyev Toshmatov", 1, "+998902222222")
    first_parent = await db.add_parent(30001, "+998903333333")
    second_parent = await db.add_parent(30002, "+998904444444")
    await db.link_student_parent(first_student, first_parent)
    await db.link_student_parent(second_student, second_parent)

    audience = await db.get_broadcast_audience("student", str(first_student))

    assert [recipient["telegram_id"] for recipient in audience] == [30001]
    assert await db.get_broadcast_audience("student", "Ali Valiyev") == []

    state = AsyncMock()
    event_order = []
    callback = SimpleNamespace(
        from_user=SimpleNamespace(id=settings.admin_ids_list[0]),
        data="bc_aud:student",
        message=SimpleNamespace(edit_text=AsyncMock(side_effect=lambda *args, **kwargs: event_order.append("edit"))),
        answer=AsyncMock(side_effect=lambda *args, **kwargs: event_order.append("answer")),
    )
    await cb_broadcast_select_audience(callback, state, db)
    assert event_order[0] == "answer"
    keyboard = callback.message.edit_text.await_args.kwargs["reply_markup"]
    button_callbacks = [button.callback_data for row in keyboard.inline_keyboard for button in row]
    assert f"bc_student:{first_student}:0" in button_callbacks

    student_callback = SimpleNamespace(
        from_user=SimpleNamespace(id=settings.admin_ids_list[0]),
        data=f"bc_student:{first_student}:0",
        message=SimpleNamespace(edit_text=AsyncMock()),
        answer=AsyncMock(),
    )
    await cb_broadcast_select_student(student_callback, state, db)
    state.update_data.assert_awaited_with(bc_aud_type="student", bc_aud_val=str(first_student))
    assert student_callback.answer.await_count == 1


@pytest.mark.asyncio
async def test_broadcast_ttl_cancel_and_long_photo_caption(tmp_path):
    db = Database(str(tmp_path / "broadcast_ttl.db"))
    await db.init_db()
    admin_id = settings.admin_ids_list[0]
    driver_tg_id = 30003
    await db.add_driver(driver_tg_id, "Driver", "+998905555555")

    PENDING_BROADCASTS.clear()
    PENDING_BROADCASTS["expired"] = {"created_at": 100.0}
    PENDING_BROADCASTS["active"] = {"created_at": 100.0}
    _purge_expired_broadcasts(now=100.0 + BROADCAST_TTL_SECONDS)
    assert "expired" not in PENDING_BROADCASTS

    PENDING_BROADCASTS["cancel-me"] = {"created_at": 100.0}
    cancel_callback = SimpleNamespace(
        from_user=SimpleNamespace(id=admin_id),
        data="bc_cancel_batch:cancel-me",
        message=SimpleNamespace(photo=[], edit_text=AsyncMock()),
        answer=AsyncMock(),
    )
    state = AsyncMock()
    await cb_broadcast_cancel(cancel_callback, state, db)
    assert "cancel-me" not in PENDING_BROADCASTS
    assert cancel_callback.answer.await_count == 1

    source_state = AsyncMock()
    source_state.get_data.return_value = {"bc_aud_type": "drivers", "bc_aud_val": None}
    source_message = SimpleNamespace(
        from_user=SimpleNamespace(id=admin_id),
        text="x" * 1100,
        caption=None,
        photo=[SimpleNamespace(file_id="photo-id")],
        answer=AsyncMock(),
        answer_photo=AsyncMock(),
    )
    await process_broadcast_content(source_message, source_state, db)
    source_message.answer.assert_awaited_once()
    source_message.answer_photo.assert_awaited_once_with(photo="photo-id")
    batch_id = next(reversed(PENDING_BROADCASTS))

    bot = AsyncMock()
    callback = SimpleNamespace(
        from_user=SimpleNamespace(id=admin_id),
        data=f"bc_send:{batch_id}",
        message=SimpleNamespace(photo=[], edit_text=AsyncMock()),
        answer=AsyncMock(),
        bot=bot,
    )
    await cb_broadcast_send(callback, db)
    bot.send_message.assert_awaited_once_with(chat_id=driver_tg_id, text="x" * 1100, parse_mode=None)
    bot.send_photo.assert_awaited_once_with(chat_id=driver_tg_id, photo="photo-id")

    stale_id = "stale001"
    await db.save_workflow_payload(stale_id, admin_id, "broadcast", {
        "created_at": time.time() - BROADCAST_TTL_SECONDS - 1,
        "audience": [{"telegram_id": driver_tg_id}],
        "text": "stale",
        "media_id": None,
        "audience_type": "drivers",
        "audience_value": None,
    })
    stale_bot = AsyncMock()
    stale_callback = SimpleNamespace(
        from_user=SimpleNamespace(id=admin_id),
        data=f"bc_send:{stale_id}",
        message=SimpleNamespace(photo=[], edit_text=AsyncMock()),
        answer=AsyncMock(),
        bot=stale_bot,
    )
    await cb_broadcast_send(stale_callback, db)
    stale_bot.send_message.assert_not_awaited()
    assert "eskirgan" in stale_callback.message.edit_text.await_args.kwargs["text"].lower()
