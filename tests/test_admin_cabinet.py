import pytest
import os
import uuid
import openpyxl
from datetime import datetime
from unittest.mock import AsyncMock, MagicMock
from aiogram.types import Message, User, CallbackQuery, Document, ReplyKeyboardMarkup

from database.db import Database
from handlers.admin import (
    get_admin_menu_kb,
    admin_today_trips_menu,
    admin_statistics_menu,
    admin_excel_menu,
    cb_admin_excel_export,
    cb_admin_excel_template,
    cb_admin_excel_upload_info,
    admin_users_menu,
    admin_today_absences,
    process_admin_excel_upload,
    cb_excel_save,
    PENDING_EXCEL_IMPORTS,
)
from services.excel_service import ExcelService
from config import settings

@pytest.mark.asyncio
async def test_admin_cabinet_full_flow():
    test_db = f"test_admin_cab_{uuid.uuid4().hex[:8]}.db"
    db = Database(test_db)

    try:
        await db.init_db()

        # 1. Admin Reply klaviaturasi tekshiruvi
        kb = get_admin_menu_kb()
        assert isinstance(kb, ReplyKeyboardMarkup)
        buttons = [b.text for row in kb.keyboard for b in row]
        expected_buttons = [
            "🚐 Bugungi reyslar", "📊 Statistika / Hisobot",
            "📥 Excel boshqaruvi", "📢 Xabar yuborish",
            "👥 Foydalanuvchilar", "🚫 Bugun kelmaganlar", "🧩 Bugungi reyslarni tuzish"
        ]
        for eb in expected_buttons:
            assert eb in buttons, f"Tugma topilmadi: {eb}"

        admin_id = settings.admin_ids_list[0]
        mock_bot = AsyncMock()

        # 2. Boshlang'ich ma'lumotlar kiritish: haydovchi, o'quvchilar, reys va absences
        driver_id = await db.add_driver(12345678, "Ali Haydovchi", "+998901112233")
        s1_id = await db.add_student("Temur Po'latov", 1, "90 555 11 22")
        s2_id = await db.add_student("Jasur Shokirov", 2, "90 555 33 44")

        # Dars jadvallari
        await db.add_schedule(s1_id, 0, "08:00", "10:00")
        await db.add_schedule(s1_id, 2, "08:00", "10:00")
        await db.add_schedule(s2_id, 1, "14:00", "16:00")

        today_str = datetime.now().strftime("%Y-%m-%d")

        # Reys yaratish (to_center)
        conn = await db.get_connection()
        cur_trip = await conn.execute("""
            INSERT INTO trips (date, direction, driver_id, planned_time, status, approved)
            VALUES (?, 'to_center', ?, '07:30', 'on_way', 0)
        """, (today_str, driver_id))
        trip_id = cur_trip.lastrowid
        await conn.execute("""
            INSERT INTO trip_items (trip_id, student_id, village_id, status)
            VALUES (?, ?, 1, 'arrived')
        """, (trip_id, s1_id))
        await conn.commit()
        await conn.close()

        # Absences qayd etish
        await db.mark_student_absent_with_reason(
            student_id=s2_id,
            date_str=today_str,
            role="parent",
            actor_id=99999,
            reason="Mazasi yo'q, shamollagan"
        )

        # 3. "📊 Statistika / Hisobot" tekshiruvi
        msg_stats = MagicMock(spec=Message)
        msg_stats.from_user = User(id=admin_id, is_bot=False, first_name="Admin")
        msg_stats.answer = AsyncMock()

        await admin_statistics_menu(msg_stats, db)
        assert msg_stats.answer.called
        stats_text = msg_stats.answer.call_args[0][0]
        assert "Bugungi transport hisoboti" in stats_text
        assert "Jami reyslar" in stats_text
        assert "Markazga yetib kelganlar" in stats_text
        assert "Bugun darsga kelmaganlar" in stats_text

        # 4. "🚫 Bugun kelmaganlar" tekshiruvi
        msg_abs = MagicMock(spec=Message)
        msg_abs.from_user = User(id=admin_id, is_bot=False, first_name="Admin")
        msg_abs.answer = AsyncMock()

        await admin_today_absences(msg_abs, db)
        assert msg_abs.answer.called
        abs_text = msg_abs.answer.call_args[0][0]
        assert "Jasur Shokirov" in abs_text
        assert "Mazasi yo'q, shamollagan" in abs_text

        # 5. "🚐 Bugungi reyslar" tekshiruvi
        msg_trips = MagicMock(spec=Message)
        msg_trips.from_user = User(id=admin_id, is_bot=False, first_name="Admin")
        msg_trips.answer = AsyncMock()

        await admin_today_trips_menu(msg_trips, db)
        assert msg_trips.answer.call_count >= 2  # Sarlavha + reys kartasi
        assert any("Temur Po'latov" in call.args[0] for call in msg_trips.answer.await_args_list)

        # 6. "👥 Foydalanuvchilar" tekshiruvi
        msg_users = MagicMock(spec=Message)
        msg_users.from_user = User(id=admin_id, is_bot=False, first_name="Admin")
        msg_users.answer = AsyncMock()

        await admin_users_menu(msg_users)
        assert msg_users.answer.called
        users_kb = msg_users.answer.call_args[1]["reply_markup"]
        btn_texts = [b.text for row in users_kb.inline_keyboard for b in row]
        assert any("Haydovchilar" in t for t in btn_texts)
        assert any("Xodimlar" in t for t in btn_texts)
        assert any("Mahallalar" in t for t in btn_texts)
        assert any("O'quvchilar" in t for t in btn_texts)

        # 7. "📥 Excel boshqaruvi" va Bazani Excelga eksport qilish
        msg_excel = MagicMock(spec=Message)
        msg_excel.from_user = User(id=admin_id, is_bot=False, first_name="Admin")
        msg_excel.answer = AsyncMock()

        await admin_excel_menu(msg_excel)
        assert msg_excel.answer.called

        # Export callback tekshiruvi
        cb_export = MagicMock(spec=CallbackQuery)
        cb_export.from_user = User(id=admin_id, is_bot=False, first_name="Admin")
        cb_export.message = MagicMock(spec=Message)
        cb_export.message.answer_document = AsyncMock()
        cb_export.answer = AsyncMock()

        await cb_admin_excel_export(cb_export, db)
        assert cb_export.message.answer_document.called
        doc_arg = cb_export.message.answer_document.call_args[0][0]
        assert doc_arg.filename == "baza_oquvchilar.xlsx"

        # ExcelService.export_students_to_excel faylini to'g'ridan-to'g'ri tekshirish
        export_file = f"test_export_{uuid.uuid4().hex[:8]}.xlsx"
        res_path = await ExcelService.export_students_to_excel(db, export_file)
        assert os.path.exists(res_path)

        wb = openpyxl.load_workbook(res_path)
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "F.I.Sh"
        # 1-o'quvchi ma'lumoti
        found_temur = False
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == "Temur Po'latov":
                found_temur = True
                assert ws.cell(row=row, column=3).value == "+998905551122"
                assert ws.cell(row=row, column=5).value == "08:00-10:00"  # Dushanba
        assert found_temur
        wb.close()
        if os.path.exists(res_path):
            os.remove(res_path)

        # 8. To'g'ridan-to'g'ri Excel fayl yuklash (Upload & Preview & Save)
        upload_xlsx = f"test_upload_{uuid.uuid4().hex[:8]}.xlsx"
        ExcelService.generate_template(upload_xlsx)
        wb_up = openpyxl.load_workbook(upload_xlsx)
        ws_up = wb_up.active
        ws_up.append([
            "Nodir Saidov", "Arpaariq", "+998931112233", "", "10:00-12:00", "", "", "", "", ""
        ])
        wb_up.save(upload_xlsx)
        wb_up.close()

        # Mock download va document
        doc_msg = MagicMock(spec=Message)
        doc_msg.from_user = User(id=admin_id, is_bot=False, first_name="Admin")
        doc_msg.document = MagicMock(spec=Document)
        doc_msg.document.file_name = "yangilar.xlsx"
        doc_msg.answer = AsyncMock()

        async def fake_download(doc, destination):
            import shutil
            shutil.copyfile(upload_xlsx, destination)

        doc_msg.bot = MagicMock()
        doc_msg.bot.download = AsyncMock(side_effect=fake_download)

        await process_admin_excel_upload(doc_msg, db)
        assert doc_msg.answer.called
        preview_text = doc_msg.answer.call_args[0][0]
        assert "Fayl tahlili" in preview_text
        assert "Yangi o'quvchilar" in preview_text

        # Saqlash tugmasini bosish
        preview_kb = doc_msg.answer.call_args[1]["reply_markup"]
        save_cb_data = preview_kb.inline_keyboard[0][0].callback_data
        assert save_cb_data.startswith("excel_save:")
        batch_id = save_cb_data.split(":")[1]

        save_cb = MagicMock(spec=CallbackQuery)
        save_cb.from_user = User(id=admin_id, is_bot=False, first_name="Admin")
        save_cb.data = save_cb_data
        save_cb.message = MagicMock(spec=Message)
        save_cb.message.edit_text = AsyncMock()
        save_cb.answer = AsyncMock()

        await cb_excel_save(save_cb, db)
        assert save_cb.answer.called

        # Yangi o'quvchi bazada paydo bo'lganini tekshirish
        cur_check = await db.get_active_students()
        names = [s["full_name"] for s in cur_check]
        assert "Nodir Saidov" in names

        if os.path.exists(upload_xlsx):
            os.remove(upload_xlsx)

    finally:
        if os.path.exists(test_db):
            os.remove(test_db)
